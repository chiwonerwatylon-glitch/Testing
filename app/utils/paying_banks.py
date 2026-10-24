"""Excel template and import for the admin-managed paying banks list."""

from __future__ import annotations

import io
import re
from typing import List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = ["Bank Name", "Account Name", "Account Number", "Currency Code"]


def build_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Paying Banks"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCD", (28, 32, 22, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().lstrip("'")


def parse_upload(file_obj) -> Tuple[List[dict], List[str]]:
    """Parse an uploaded workbook. Returns (rows, errors)."""
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        return [], ["The file is not a valid Excel (.xlsx) workbook."]
    ws = wb.active
    rows, errors = [], []
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header or [_text(h).lower() for h in header[:4]] != [h.lower() for h in HEADERS]:
        return [], [
            "The first row must be the headings: " + ", ".join(HEADERS) + ". Download the template."
        ]
    for n, r in enumerate(it, start=2):
        r = list(r) + [None] * 4
        bank, name, number, cur = (_text(v) for v in r[:4])
        if not (bank or name or number or cur):
            continue
        cur = cur.upper()
        if not (bank and name and number and cur):
            errors.append(f"Row {n}: all four columns are required.")
        elif not re.fullmatch(r"[A-Z]{3}", cur):
            errors.append(f"Row {n}: currency code '{cur}' must be 3 letters.")
        else:
            rows.append(
                {"bank_name": bank, "account_name": name, "account_number": number, "currency_code": cur}
            )
    if not rows and not errors:
        errors.append("The file has no data rows.")
    return rows, errors
