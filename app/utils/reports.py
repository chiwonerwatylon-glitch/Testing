"""Report builders for requisition CSV and Excel exports."""
import csv
import io
import os

from flask import current_app

BRAND_RED = 'D21F3C'
BRAND_BLACK = '1A1A1A'

COMPANY_NAME = 'First Mutual Life Assurance Company'
AMOUNT_INDEX = 7

COLUMNS = [
    ('Requisition No', lambda r: r.requisition_number),
    ('Status', lambda r: r.status.replace('_', ' ').title()),
    ('Payee', lambda r: r.in_favor_of or ''),
    ('Bank', lambda r: r.bank_name or ''),
    ('Branch', lambda r: r.branch or ''),
    ('Account No', lambda r: r.account_number or ''),
    ('Currency', lambda r: r.currency or ''),
    ('Amount', lambda r: f"{(r.total_debit or 0):.2f}"),
    ('Initiator', lambda r: r.initiator.full_name if r.initiator else ''),
    ('Reviewer', lambda r: r.reviewer.full_name if r.reviewer else ''),
    ('Approvals', lambda r: f"{r.current_approvals or 0}/{r.required_approvals or 0}"),
    ('Departments', lambda r: ', '.join(sorted({i.account_department for i in r.items if i.account_department}))),
    ('Created', lambda r: r.created_at.strftime('%Y-%m-%d %H:%M') if r.created_at else ''),
    ('Completed', lambda r: r.completed_at.strftime('%Y-%m-%d %H:%M') if r.completed_at else ''),
]


def logo_path():
    return os.path.join(current_app.root_path, 'static', 'images', 'logo.png')


def build_rows(requisitions):
    headers = [header for header, _getter in COLUMNS]
    rows = [[getter(r) for _header, getter in COLUMNS] for r in requisitions]
    return headers, rows


def totals_by_currency(requisitions):
    totals = {}
    for r in requisitions:
        code = r.currency or ''
        totals[code] = totals.get(code, 0) + (r.total_debit or 0)
    if not totals:
        return '-'
    return ' | '.join(f"{code} {totals[code]:,.2f}" for code in sorted(totals))


def describe_filters(filters):
    parts = []
    if filters.get('status'):
        parts.append(f"Status: {filters['status']}")
    if filters.get('currency'):
        parts.append(f"Currency: {filters['currency']}")
    if filters.get('department'):
        parts.append(f"Department: {filters['department']}")
    if filters.get('initiator_name'):
        parts.append(f"Initiator: {filters['initiator_name']}")
    if filters.get('date_from'):
        parts.append(f"From: {filters['date_from']}")
    if filters.get('date_to'):
        parts.append(f"To: {filters['date_to']}")
    return '; '.join(parts) if parts else 'All records'


def to_csv(requisitions, meta):
    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow([COMPANY_NAME])
    writer.writerow(['Requisition Report'])
    writer.writerow([f"Generated: {meta['generated']}"])
    writer.writerow([f"Filters: {meta['filters']}"])
    writer.writerow([])
    headers, rows = build_rows(requisitions)
    writer.writerow(headers)
    writer.writerows(rows)
    buf = io.BytesIO(('\ufeff' + text.getvalue()).encode('utf-8'))
    buf.seek(0)
    return buf


def to_excel(requisitions, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Requisitions'

    # Logo (optional)
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(logo_path())
        ratio = img.width / img.height
        img.height = 60
        img.width = int(60 * ratio)
        ws.add_image(img, 'A1')
    except Exception:
        pass

    ws['C1'] = COMPANY_NAME
    ws['C1'].font = Font(bold=True, size=16, color=BRAND_RED)
    ws['C2'] = 'Requisition Report'
    ws['C2'].font = Font(bold=True, size=12)
    ws['C3'] = f"Generated: {meta['generated']}"
    ws['C3'].font = Font(size=9, color='666666')
    ws['C4'] = f"Filters: {meta['filters']}"
    ws['C4'].font = Font(size=9, color='666666')

    side = Side(style='thin', color='D0D0D0')
    border = Border(left=side, right=side, top=side, bottom=side)

    headers, rows = build_rows(requisitions)
    header_row = 6
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=BRAND_BLACK, end_color=BRAND_BLACK, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(header) + 4)

    first_data = header_row + 1
    row_num = first_data
    for row in rows:
        for col, value in enumerate(row, start=1):
            if col - 1 == AMOUNT_INDEX:
                # Write a real number so the SUM total works
                value = float(value)
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = border
            if col - 1 == AMOUNT_INDEX:
                cell.number_format = '#,##0.00'
        row_num += 1

    ws.freeze_panes = 'A7'

    # Totals row
    last = row_num - 1
    amount_letter = get_column_letter(AMOUNT_INDEX + 1)
    ws.cell(row=row_num, column=1, value='TOTAL').font = Font(bold=True)
    total_cell = ws.cell(row=row_num, column=AMOUNT_INDEX + 1,
                         value=f"=SUM({amount_letter}{first_data}:{amount_letter}{last})")
    total_cell.font = Font(bold=True)
    total_cell.number_format = '#,##0.00'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
