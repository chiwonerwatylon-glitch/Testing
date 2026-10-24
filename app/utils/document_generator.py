from __future__ import annotations

import csv
import io
import os
from datetime import datetime

from flask import current_app
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from xml.sax.saxutils import escape as _xml_escape

# FML Brand Colors
FML_RED = colors.HexColor("#D21F3C")
FML_BLACK = colors.HexColor("#1A1A1A")
FML_RED_HEX = "D21F3C"
FML_BLACK_HEX = "1A1A1A"


class _LinkedQR(Image):
    """QR code image that is also a clickable link in the PDF."""

    def __init__(self, png_bytes, url, size):
        super().__init__(io.BytesIO(png_bytes), width=size, height=size)
        self._url = url

    def draw(self):
        super().draw()
        if self._url:
            self.canv.linkURL(
                self._url, (0, 0, self.drawWidth, self.drawHeight), relative=1
            )


class RequisitionDocumentGenerator:
    """Generate PDF, DOCX, Excel and CSV representations of a requisition."""

    def __init__(self, requisition):
        self.requisition = requisition
        self.logo_path = os.path.join(
            current_app.root_path, "static", "images", "logo.png"
        )

    def _paying_bank_text(self) -> str:
        r = self.requisition
        if not r.paying_bank_name:
            return ""
        parts = [r.paying_bank_name, r.paying_account_name, r.paying_account_number]
        text = " - ".join(p for p in parts if p)
        extra = [x for x in (r.paying_currency, r.paying_swift_code and f"SWIFT {r.paying_swift_code}") if x]
        return f"{text} ({', '.join(extra)})" if extra else text

    # ---------------------------------------------------------------------
    # PDF
    # ---------------------------------------------------------------------

    def _payment_stamp(self):
        """(text, hex colour) for fully approved requisitions, else None."""
        if getattr(self.requisition, "status", None) != "approved":
            return None
        if getattr(self.requisition, "is_paid", False):
            return "PAID", "198754"
        return "NOT PAID", "DC3545"

    def generate_pdf(self) -> io.BytesIO:
        """Generate PDF version of requisition."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=1 * cm,
            leftMargin=1 * cm,
            topMargin=1 * cm,
            bottomMargin=1 * cm,
        )

        styles = getSampleStyleSheet()
        elements = []

        # Custom styles - FML Red & Black theme
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=FML_BLACK,
            alignment=TA_CENTER,
            spaceAfter=10,
        )

        subtitle_style = ParagraphStyle(
            "Subtitle",
            parent=styles["Normal"],
            fontSize=12,
            textColor=FML_RED,
            alignment=TA_CENTER,
            spaceAfter=10,
            fontName="Helvetica-Bold",
        )

        ref_style = ParagraphStyle(
            "RefStyle",
            parent=styles["Normal"],
            fontSize=10,
            alignment=TA_CENTER,
            spaceAfter=15,
        )

        # Header with logo
        if os.path.exists(self.logo_path):
            logo = Image(self.logo_path, width=2 * inch, height=0.8 * inch)
            elements.append(logo)

        elements.append(
            Paragraph("FIRST MUTUAL LIFE ASSURANCE COMPANY", title_style)
        )
        elements.append(Paragraph("PAYMENT REQUISITION", subtitle_style))
        elements.append(
            Paragraph(
                f"Requisition No: {self.requisition.requisition_number}", ref_style
            )
        )

        stamp = self._payment_stamp()
        if stamp:
            stamp_color = colors.HexColor("#" + stamp[1])
            stamp_table = Table(
                [[Paragraph(
                    f'<font color="#{stamp[1]}"><b>{stamp[0]}</b></font>',
                    ParagraphStyle(
                        "Stamp", parent=styles["Normal"], fontSize=22,
                        alignment=TA_CENTER, leading=26,
                    ),
                )]],
                colWidths=[2.4 * inch],
                style=TableStyle([
                    ("BOX", (0, 0), (-1, -1), 3, stamp_color),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]),
            )
            elements.append(stamp_table)
            elements.append(Spacer(1, 0.15 * inch))

        # Red divider line
        elements.append(
            Table(
                [[""]],
                colWidths=[7 * inch],
                style=TableStyle(
                    [("LINEBELOW", (0, 0), (-1, -1), 2, FML_RED)]
                ),
            )
        )
        elements.append(Spacer(1, 0.3 * inch))

        # Payment details: label/value pairs side by side
        label_bg = colors.HexColor("#F8F8F8")
        cell_style = ParagraphStyle(
            "DetailCell", parent=styles["Normal"], fontSize=9.5, leading=12
        )
        bold_style = ParagraphStyle(
            "DetailLabel", parent=cell_style, fontName="Helvetica-Bold"
        )

        def _p(text, style=cell_style):
            return Paragraph(_xml_escape(str(text or "")), style)

        rows = [[_p("Paying Bank:", bold_style), _p(self._paying_bank_text()), "", ""]]
        for a, b, c, d in self._detail_rows():
            rows.append([_p(a, bold_style), _p(b), _p(c, bold_style), _p(d)])
        rows.append([_p("Reason for Payment:", bold_style), _p(self.requisition.reason_for_payment), "", ""])

        last = len(rows) - 1
        payment_table = Table(
            rows, colWidths=[1.25 * inch, 2.25 * inch, 1.25 * inch, 2.25 * inch]
        )
        payment_table.setStyle(
            TableStyle(
                [
                    ("SPAN", (1, 0), (3, 0)),
                    ("SPAN", (1, last), (3, last)),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
                    ("BACKGROUND", (0, 0), (0, -1), label_bg),
                    ("BACKGROUND", (2, 1), (2, last - 1), label_bg),
                ]
            )
        )
        elements.append(payment_table)
        elements.append(Spacer(1, 0.3 * inch))

        # Items table
        items_header = [
            "Account Name",
            "Account Code",
            "Location",
            "Department",
            "Details",
            "Debit",
            "Credit",
        ]
        items_data = [items_header]

        for item in self.requisition.items:
            items_data.append(
                [
                    item.account_name,
                    item.account_code,
                    item.account_location or "",
                    item.account_department or "",
                    item.details or "",
                    f"${item.debit:,.2f}" if item.debit else "",
                    f"${item.credit:,.2f}" if item.credit else "",
                ]
            )

        # Totals row
        items_data.append(
            [
                "",
                "",
                "",
                "",
                "TOTALS",
                f"${self.requisition.total_debit:,.2f}",
                f"${self.requisition.total_credit:,.2f}",
            ]
        )

        items_table = Table(
            items_data,
            colWidths=[
                1.2 * inch,
                0.9 * inch,
                0.8 * inch,
                0.9 * inch,
                1.3 * inch,
                0.9 * inch,
                0.9 * inch,
            ],
        )
        items_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), FML_BLACK),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (5, 0), (-1, -1), "RIGHT"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    (
                        "BACKGROUND",
                        (0, -1),
                        (-1, -1),
                        colors.HexColor("#FCE8EA"),
                    ),
                    ("TEXTCOLOR", (0, -1), (-1, -1), FML_RED),
                ]
            )
        )
        elements.append(items_table)
        elements.append(Spacer(1, 0.25 * inch))

        # Signatures section
        auth_style = ParagraphStyle(
            "AuthStyle",
            parent=styles["Heading3"],
            textColor=FML_RED,
        )
        elements.append(Paragraph("<b>AUTHORIZATION</b>", auth_style))
        elements.append(Spacer(1, 0.2 * inch))

        sig_name_style = ParagraphStyle(
            "SigName", parent=styles["Normal"], fontSize=9, leading=11
        )
        sig_data = [["Role", "Name & Signature", "Date & Time Signed", "Signature Hash"]]
        for sg in self._signers():
            name_cell = [Paragraph(f"<b>{_xml_escape(sg['name'])}</b>", sig_name_style)]
            if sg["sig"]:
                w, h = self._fit(sg["sig"], 1.3 * inch, 0.42 * inch)
                name_cell.append(Spacer(1, 3))
                name_cell.append(Image(sg["sig"], width=w, height=h, hAlign="LEFT"))
            sig_data.append([sg["role"], name_cell, sg["when"], sg["hash"]])

        sig_table = Table(
            sig_data,
            colWidths=[1.3 * inch, 2.2 * inch, 1.6 * inch, 1.9 * inch],
        )
        sig_table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("BACKGROUND", (0, 0), (-1, 0), FML_BLACK),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ]
            )
        )
        elements.append(sig_table)

        # Document security box
        security = self._get_security_info()
        value_style = ParagraphStyle(
            "SecurityValue",
            parent=styles["Normal"],
            fontSize=9,
            fontName="Helvetica",
            leading=11,
        )
        label_style = ParagraphStyle(
            "SecurityLabel",
            parent=value_style,
            fontName="Helvetica-Bold",
        )
        note_style = ParagraphStyle(
            "SecurityNote",
            parent=value_style,
            fontName="Helvetica-Oblique",
            fontSize=8,
        )

        def _cell(value: str) -> Paragraph:
            return Paragraph(_xml_escape(str(value or "-")), value_style)

        security_data = [
            [
                Paragraph("Document code", label_style),
                _cell(security.get("document_code")),
            ],
            [
                Paragraph("Authorised amount", label_style),
                _cell(security.get("authorised_amount")),
            ],
            [
                Paragraph("Verify at", label_style),
                Paragraph(
                    f'<a href="{_xml_escape(str(security.get("verify_url") or ""))}" color="blue">'
                    f'{_xml_escape(str(security.get("verify_url") or "-"))}</a>',
                    value_style,
                ),
            ],
            [
                Paragraph(
                    "Any change to amounts, status or approvals invalidates this "
                    "document code.",
                    note_style,
                ),
                "",
            ],
        ]
        security_table = Table(
            security_data,
            colWidths=[1.4 * inch, 3.6 * inch],
        )
        security_table.setStyle(
            TableStyle(
                [
                    ("BOX", (0, 0), (-1, -1), 1, FML_BLACK),
                    ("SPAN", (0, -1), (-1, -1)),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BACKGROUND", (0, 0), (0, -2), colors.HexColor("#F8F8F8")),
                ]
            )
        )
        qr_png = self._qr_png(security.get("verify_url") or "")
        if qr_png:
            qr_cell = [
                _LinkedQR(qr_png, security.get("verify_url"), 1.25 * inch),
                Paragraph(
                    f'<a href="{_xml_escape(str(security.get("verify_url")))}" color="blue">Scan or click to verify</a>',
                    ParagraphStyle("QRCap", parent=note_style, alignment=TA_CENTER),
                ),
            ]
            security_block = Table(
                [[security_table, qr_cell]],
                colWidths=[5.1 * inch, 1.9 * inch],
                style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0)]),
            )
        else:
            security_block = security_table
        elements.append(Spacer(1, 0.3 * inch))
        elements.append(
            KeepTogether(
                [
                    Paragraph("<b>DOCUMENT SECURITY</b>", auth_style),
                    Spacer(1, 0.1 * inch),
                    security_block,
                ]
            )
        )

        # Footer
        elements.append(Spacer(1, 0.3 * inch))
        footer_style = ParagraphStyle(
            "Footer",
            parent=styles["Normal"],
            fontSize=8,
            textColor=colors.grey,
            alignment=TA_CENTER,
        )
        elements.append(
            Paragraph(
                "This is a system-generated document from First Mutual Life "
                "Assurance Company Requisition System",
                footer_style,
            )
        )

        doc.build(elements)
        buffer.seek(0)
        return buffer

    # ---------------------------------------------------------------------
    # DOCX
    # ---------------------------------------------------------------------

    def generate_docx(self) -> io.BytesIO:
        """Generate DOCX version of requisition."""
        doc = Document()

        # Set margins
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)

        # Header
        header_para = doc.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if os.path.exists(self.logo_path):
            run = header_para.add_run()
            run.add_picture(self.logo_path, width=Inches(2))

        title = doc.add_paragraph("FIRST MUTUAL LIFE ASSURANCE COMPANY")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.runs[0]
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

        subtitle = doc.add_paragraph("PAYMENT REQUISITION")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.runs[0]
        subtitle_run.bold = True
        subtitle_run.font.size = Pt(14)
        subtitle_run.font.color.rgb = RGBColor(0xD2, 0x1F, 0x3C)

        req_num = doc.add_paragraph(
            f"Requisition No: {self.requisition.requisition_number}"
        )
        req_num.alignment = WD_ALIGN_PARAGRAPH.CENTER

        stamp = self._payment_stamp()
        if stamp:
            stamp_p = doc.add_paragraph()
            stamp_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            stamp_run = stamp_p.add_run(f"  {stamp[0]}  ")
            stamp_run.bold = True
            stamp_run.font.size = Pt(24)
            stamp_run.font.color.rgb = RGBColor.from_string(stamp[1])
            rPr = stamp_run._r.get_or_add_rPr()
            bdr = OxmlElement("w:bdr")
            for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "4"), ("w:color", stamp[1])):
                bdr.set(qn(k), v)
            rPr.append(bdr)

        doc.add_paragraph()

        # Payment details table
        detail_rows = self._detail_rows()
        payment_table = doc.add_table(rows=len(detail_rows) + 2, cols=4)
        payment_table.style = "Table Grid"

        def _set(cell, text, bold=False):
            cell.text = str(text or "")
            if bold and cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].bold = True

        top = payment_table.rows[0].cells
        _set(top[0], "Paying Bank:", True)
        _set(top[1], self._paying_bank_text())
        top[1].merge(top[3])
        for i, (a, b, c, d) in enumerate(detail_rows, 1):
            cells = payment_table.rows[i].cells
            _set(cells[0], a, True)
            _set(cells[1], b)
            _set(cells[2], c, True)
            _set(cells[3], d)
        bottom = payment_table.rows[-1].cells
        _set(bottom[0], "Reason for Payment:", True)
        _set(bottom[1], self.requisition.reason_for_payment)
        bottom[1].merge(bottom[3])

        doc.add_paragraph()

        # Items table
        items = list(self.requisition.items)
        items_table = doc.add_table(rows=len(items) + 2, cols=7)
        items_table.style = "Table Grid"

        headers = [
            "Account Name",
            "Account Code",
            "Location",
            "Department",
            "Details",
            "Debit",
            "Credit",
        ]
        header_row = items_table.rows[0]
        for i, header in enumerate(headers):
            cell = header_row.cells[i]
            cell.text = header
            run = cell.paragraphs[0].runs[0]
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # Set black background for header cell
            shading_elm = cell._element
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), FML_BLACK_HEX)
            shading_elm.get_or_add_tcPr().append(shd)

        for i, item in enumerate(items, 1):
            row = items_table.rows[i]
            row.cells[0].text = item.account_name
            row.cells[1].text = item.account_code
            row.cells[2].text = item.account_location or ""
            row.cells[3].text = item.account_department or ""
            row.cells[4].text = item.details or ""
            row.cells[5].text = (
                f"${item.debit:,.2f}" if item.debit else ""
            )
            row.cells[6].text = (
                f"${item.credit:,.2f}" if item.credit else ""
            )

        # Totals row
        totals_row = items_table.rows[-1]
        totals_row.cells[4].text = "TOTALS"
        totals_row.cells[5].text = f"${self.requisition.total_debit:,.2f}"
        totals_row.cells[6].text = f"${self.requisition.total_credit:,.2f}"
        for cell in totals_row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xD2, 0x1F, 0x3C)

        doc.add_paragraph()
        auth_para = doc.add_paragraph("AUTHORIZATION")
        auth_run = auth_para.runs[0]
        auth_run.bold = True
        auth_run.font.color.rgb = RGBColor(0xD2, 0x1F, 0x3C)
        auth_run.font.size = Pt(13)

        # Signatures
        self._add_docx_signatures(doc)

        # Document security
        doc.add_paragraph()
        security_para = doc.add_paragraph("DOCUMENT SECURITY")
        sec_run = security_para.runs[0]
        sec_run.bold = True
        sec_run.font.color.rgb = RGBColor(0xD2, 0x1F, 0x3C)
        sec_run.font.size = Pt(13)

        security = self._get_security_info()
        security_rows = [
            ("Document code", security.get("document_code") or "-"),
            ("Authorised amount", security.get("authorised_amount") or "-"),
            ("Verify at", security.get("verify_url") or "-"),
        ]
        security_table = doc.add_table(rows=len(security_rows), cols=2)
        security_table.style = "Table Grid"
        for i, (label, value) in enumerate(security_rows):
            security_table.rows[i].cells[0].text = label
            security_table.rows[i].cells[1].text = str(value)
            cell_run = security_table.rows[i].cells[0].paragraphs[0].runs[0]
            cell_run.bold = True

        verify_url = security.get("verify_url") or ""
        qr_png = self._qr_png(verify_url)
        if qr_png and verify_url:
            qr_para = doc.add_paragraph()
            qr_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = qr_para.add_run()
            run.add_picture(io.BytesIO(qr_png), width=Inches(1.3))
            # make the QR picture itself clickable
            from docx.opc.constants import RELATIONSHIP_TYPE as RT

            r_id = qr_para.part.relate_to(verify_url, RT.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), r_id)
            qr_para._p.remove(run._r)
            link.append(run._r)
            qr_para._p.append(link)
            cap = doc.add_paragraph()
            self._docx_link(cap, verify_url, "Scan or click to verify")

        note_para = doc.add_paragraph()
        note_run = note_para.add_run(
            "Any change to amounts, status or approvals invalidates this document code."
        )
        note_run.italic = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    # ---------------------------------------------------------------------
    # Excel
    # ---------------------------------------------------------------------

    def generate_excel(self) -> io.BytesIO:
        """Generate Excel version of requisition."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Requisition"

        # Styles - FML Red & Black Theme
        header_font = Font(bold=True, size=14, color=FML_BLACK_HEX)
        subheader_font = Font(bold=True, size=11, color=FML_RED_HEX)
        border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )
        header_fill = PatternFill(
            start_color=FML_BLACK_HEX,
            end_color=FML_BLACK_HEX,
            fill_type="solid",
        )
        header_font_white = Font(bold=True, color="FFFFFF")
        totals_fill = PatternFill(
            start_color="FCE8EA", end_color="FCE8EA", fill_type="solid"
        )
        totals_font = Font(bold=True, color=FML_RED_HEX)

        # Title
        ws.merge_cells("A1:G1")
        ws["A1"] = "FIRST MUTUAL LIFE ASSURANCE COMPANY"
        ws["A1"].font = header_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:G2")
        ws["A2"] = "PAYMENT REQUISITION"
        ws["A2"].font = subheader_font
        ws["A2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A3:G3")
        ws["A3"] = f"Requisition No: {self.requisition.requisition_number}"
        ws["A3"].alignment = Alignment(horizontal="center")

        stamp = self._payment_stamp()
        if stamp:
            ws.merge_cells("A4:G4")
            ws["A4"] = stamp[0]
            ws["A4"].font = Font(bold=True, size=18, color=stamp[1])
            ws["A4"].alignment = Alignment(horizontal="center")

        # Payment details: label/value pairs side by side
        row = 5
        ws[f"A{row}"] = "Paying Bank:"
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"] = self._paying_bank_text()
        ws.merge_cells(f"B{row}:G{row}")
        for col in "ABCDEFG":
            ws[f"{col}{row}"].border = border
        row += 1
        for a, b, c, d in self._detail_rows():
            ws[f"A{row}"], ws[f"B{row}"], ws[f"D{row}"], ws[f"E{row}"] = a, b, c, d
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"D{row}"].font = Font(bold=True)
            ws.merge_cells(f"B{row}:C{row}")
            ws.merge_cells(f"E{row}:G{row}")
            for col in "ABCDEFG":
                ws[f"{col}{row}"].border = border
            row += 1
        ws[f"A{row}"] = "Reason for Payment:"
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"] = self.requisition.reason_for_payment
        ws.merge_cells(f"B{row}:G{row}")
        for col in "ABCDEFG":
            ws[f"{col}{row}"].border = border
        row += 2

        # Items header
        headers = [
            "Account Name",
            "Account Code",
            "Location",
            "Department",
            "Details",
            "Debit",
            "Credit",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        row += 1

        # Items
        for item in self.requisition.items:
            ws.cell(row=row, column=1, value=item.account_name).border = border
            ws.cell(row=row, column=2, value=item.account_code).border = border
            ws.cell(
                row=row, column=3, value=item.account_location or ""
            ).border = border
            ws.cell(
                row=row, column=4, value=item.account_department or ""
            ).border = border
            ws.cell(
                row=row, column=5, value=item.details or ""
            ).border = border
            ws.cell(
                row=row, column=6, value=item.debit if item.debit else ""
            ).border = border
            ws.cell(
                row=row, column=7, value=item.credit if item.credit else ""
            ).border = border
            row += 1

        # Totals
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = totals_fill
        ws.cell(row=row, column=5, value="TOTALS").font = totals_font
        total_debit_cell = ws.cell(
            row=row,
            column=6,
            value=float(self.requisition.total_debit or 0),
        )
        total_debit_cell.font = totals_font
        total_debit_cell.number_format = "#,##0.00"
        total_credit_cell = ws.cell(
            row=row,
            column=7,
            value=float(self.requisition.total_credit or 0),
        )
        total_credit_cell.font = totals_font
        total_credit_cell.number_format = "#,##0.00"

        # Column widths
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 25
        ws.column_dimensions["F"].width = 12
        ws.column_dimensions["G"].width = 12

        row += 2

        # Signatures
        ws[f"A{row}"] = "AUTHORIZATION"
        ws[f"A{row}"].font = Font(bold=True, size=12, color=FML_RED_HEX)
        row += 1

        self._add_excel_signatures(ws, row)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # ---------------------------------------------------------------------
    # CSV
    # ---------------------------------------------------------------------

    def generate_csv(self) -> io.BytesIO:
        """Generate CSV version of requisition."""
        text_buffer = io.StringIO()
        writer = csv.writer(text_buffer)

        # Header info
        writer.writerow(["FIRST MUTUAL LIFE ASSURANCE COMPANY"])
        writer.writerow(["PAYMENT REQUISITION"])
        writer.writerow(
            [f"Requisition No: {self.requisition.requisition_number}"]
        )
        stamp = self._payment_stamp()
        if stamp:
            writer.writerow([f"Payment Status: {stamp[0]}"])
        writer.writerow([])

        # Payment details
        writer.writerow(["Paying Bank:", self._paying_bank_text()])
        for row_cells in self._detail_rows():
            writer.writerow(row_cells)
        writer.writerow(["Reason for Payment:", self.requisition.reason_for_payment])
        writer.writerow([])

        # Items        writer.writerow([])

        # Items
        writer.writerow(
            [
                "Account Name",
                "Account Code",
                "Location",
                "Department",
                "Details",
                "Debit",
                "Credit",
            ]
        )
        for item in self.requisition.items:
            writer.writerow(
                [
                    item.account_name,
                    item.account_code,
                    item.account_location or "",
                    item.account_department or "",
                    item.details or "",
                    item.debit if item.debit else "",
                    item.credit if item.credit else "",
                ]
            )

        writer.writerow(
            [
                "",
                "",
                "",
                "",
                "TOTALS",
                self.requisition.total_debit,
                self.requisition.total_credit,
            ]
        )
        writer.writerow([])

        # Signatures
        writer.writerow(["AUTHORIZATION"])
        writer.writerow(["Role", "Name", "Date & Time Signed", "Signature Hash"])
        for sg in self._signers():
            writer.writerow([sg["role"], sg["name"], sg["when"], sg["hash"]])

        # Document security
        writer.writerow([])
        writer.writerow(["DOCUMENT SECURITY"])
        security = self._get_security_info()
        writer.writerow(["Document code", security["document_code"]])
        writer.writerow(["Authorised amount", security["authorised_amount"]])
        writer.writerow(["Verify at", security["verify_url"]])

        # Convert to bytes
        text_buffer.seek(0)
        bytes_buffer = io.BytesIO()
        bytes_buffer.write(text_buffer.getvalue().encode("utf-8"))
        bytes_buffer.seek(0)
        return bytes_buffer

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    # ------------------------------------------------------------------
    # Shared content: side-by-side details, signers, QR
    # ------------------------------------------------------------------

    @staticmethod
    def _fmt_long(value, default: str = "-") -> str:
        """e.g. '24 October 2026, 13:36'."""
        if not value:
            return default
        try:
            return value.strftime("%d %B %Y, %H:%M")
        except Exception:
            return default

    @staticmethod
    def _fmt_ts(value, default: str = "Pending") -> str:
        """Signature timestamp with seconds, e.g. '24 Oct 2026, 13:36:05'."""
        if not value:
            return default
        try:
            return value.strftime("%d %b %Y, %H:%M:%S")
        except Exception:
            return default

    def _detail_pairs(self):
        """(label, value) pairs shown two per row, side by side."""
        r = self.requisition
        return [
            ("In Favor Of:", r.in_favor_of or ""),
            ("Bank:", r.bank_name or ""),
            ("Branch:", r.branch or ""),
            ("Account Number:", r.account_number or ""),
            ("Currency:", r.currency or ""),
            ("ITF 263:", r.itf_263 or ""),
            ("Date Created:", self._fmt_long(getattr(r, "created_at", None))),
            ("Date Completed:", self._fmt_long(getattr(r, "completed_at", None))),
        ]

    def _detail_rows(self):
        """Rows of up to 4 cells: label, value, label, value."""
        pairs = self._detail_pairs()
        return [
            [*pairs[i], *(pairs[i + 1] if i + 1 < len(pairs) else ("", ""))]
            for i in range(0, len(pairs), 2)
        ]

    def _sig_file(self, filename):
        if not filename:
            return None
        try:
            path = os.path.join(
                current_app.root_path, "static", "signatures",
                os.path.basename(str(filename)),
            )
            return path if os.path.exists(path) else None
        except Exception:  # pragma: no cover
            return None

    def _signers(self):
        """Everyone involved, in order: role, name, signature file, timestamp, hash."""
        req = self.requisition
        out = []

        def add(role, user, sig, when, hash_text="-"):
            if user:
                out.append({
                    "role": role,
                    "name": self._user_name(user, "Pending"),
                    "sig": self._sig_file(sig) or self._sig_file(getattr(user, "signature_path", None)),
                    "when": self._fmt_ts(when),
                    "hash": hash_text,
                })
            else:
                out.append({"role": role, "name": "Pending", "sig": None,
                            "when": "Pending", "hash": "-"})

        add("Initiated By", getattr(req, "initiator", None),
            getattr(req, "initiator_signature", None),
            getattr(req, "initiated_at", None) or getattr(req, "created_at", None),
            self._short_hash(req, "initiator_signature_hash", "-"))
        add("Reviewed By", getattr(req, "reviewer", None),
            getattr(req, "reviewer_signature", None),
            getattr(req, "reviewed_at", None),
            self._short_hash(req, "reviewer_signature_hash", "-"))

        labels = {1: "1st Approver", 2: "2nd Approver", 3: "Senior Approver (3rd)"}
        logs = self._approved_logs()
        for level in range(1, (req.required_approvals or 0) + 1):
            log = next((l for l in logs if l.approval_level == level), None)
            role = labels.get(level, f"Approver {level}")
            if log:
                add(role, getattr(log, "approver", None),
                    getattr(log, "signature_path", None),
                    getattr(log, "created_at", None),
                    self._short_hash(log, "signature_hash", "-"))
            else:
                add(role, None, None, None)

        if getattr(req, "is_paid", False) and getattr(req, "paid_by", None):
            add("Marked Paid By", req.paid_by, None, getattr(req, "paid_at", None))
        return out

    def _qr_png(self, url: str):
        """PNG bytes of a QR code for the verification URL, or None."""
        try:
            import qrcode

            buf = io.BytesIO()
            qrcode.make(url).save(buf, format="PNG")
            return buf.getvalue()
        except Exception:  # pragma: no cover
            return None

    @staticmethod
    def _fit(path, max_w, max_h):
        """(width, height) scaled to fit the box, keeping the aspect ratio."""
        from PIL import Image as PILImage

        with PILImage.open(path) as im:
            w, h = im.size
        scale = min(max_w / w, max_h / h)
        return w * scale, h * scale


    def _add_docx_signatures(self, doc: Document) -> None:
        """Signatures section: everyone involved with name, signature image and timestamp."""
        signers = self._signers()
        sig_table = doc.add_table(rows=1 + len(signers), cols=4)
        sig_table.style = "Table Grid"

        headers = ["Role", "Name & Signature", "Date & Time Signed", "Signature Hash"]
        for i, header in enumerate(headers):
            cell = sig_table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), FML_BLACK_HEX)
            cell._element.get_or_add_tcPr().append(shd)

        for idx, sg in enumerate(signers, 1):
            cells = sig_table.rows[idx].cells
            cells[0].text = sg["role"]
            cells[1].text = sg["name"]
            cells[2].text = sg["when"]
            cells[3].text = sg["hash"]
            if sg["sig"]:
                try:
                    run = cells[1].paragraphs[0].add_run()
                    run.add_break()
                    run.add_picture(sg["sig"], width=Inches(1.4))
                except Exception as e:  # pragma: no cover
                    current_app.logger.warning(
                        "Could not add signature image to DOCX: %s", e
                    )

    @staticmethod
    def _docx_link(paragraph, url: str, text: str) -> None:
        """Append a clickable hyperlink run to a DOCX paragraph."""
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(color)
        rpr.append(underline)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        link.append(run)
        paragraph._p.append(link)

    def _add_excel_signatures(self, ws, start_row: int) -> int:
        """Signatures (with images and timestamps) and document security with QR."""
        from openpyxl.drawing.image import Image as XLImage

        header_fill = PatternFill(
            start_color=FML_BLACK_HEX, end_color=FML_BLACK_HEX, fill_type="solid"
        )
        header_font_white = Font(bold=True, color="FFFFFF")

        row = start_row
        headers = ["Role", "Name", "Date & Time Signed", "Signature Hash", "Signature"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 24

        for sg in self._signers():
            row += 1
            ws.cell(row=row, column=1, value=sg["role"])
            ws.cell(row=row, column=2, value=sg["name"])
            ws.cell(row=row, column=3, value=sg["when"])
            ws.cell(row=row, column=4, value=sg["hash"])
            if sg["sig"]:
                try:
                    w, h = self._fit(sg["sig"], 140, 50)
                    img = XLImage(sg["sig"])
                    img.width, img.height = int(w), int(h)
                    ws.add_image(img, f"E{row}")
                    ws.row_dimensions[row].height = max(40, h * 0.78 + 4)
                except Exception as e:  # pragma: no cover
                    current_app.logger.warning("Could not add signature image: %s", e)

        row += 2
        ws.cell(row=row, column=1, value="DOCUMENT SECURITY").font = Font(
            bold=True, size=12, color=FML_RED_HEX
        )
        row += 1
        security = self._get_security_info()
        for label, value in (
            ("Document code", security["document_code"]),
            ("Authorised amount", security["authorised_amount"]),
            ("Verify at", security["verify_url"]),
        ):
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            cell = ws.cell(row=row, column=2, value=value)
            if label == "Verify at" and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            row += 1

        qr_png = self._qr_png(security["verify_url"])
        if qr_png:
            try:
                qr = XLImage(io.BytesIO(qr_png))
                qr.width = qr.height = 110
                ws.add_image(qr, f"A{row + 1}")
                row += 8
                link = ws.cell(row=row, column=1, value="Scan or click to verify")
                link.hyperlink = security["verify_url"]
                link.font = Font(color="0563C1", underline="single")
                row += 1
            except Exception as e:  # pragma: no cover
                current_app.logger.warning("Could not add QR code: %s", e)

        return row

    # ------------------------------------------------------------------
    # Shared helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _format_dt(value, default: str = "") -> str:
        """Format a datetime as '%Y-%m-%d %H:%M', or return default."""
        if not value:
            return default
        try:
            return value.strftime("%Y-%m-%d %H:%M")
        except Exception:
            return default

    @staticmethod
    def _user_name(user, default: str = "") -> str:
        """Full name of a user, or default if the user/name is missing."""
        return getattr(user, "full_name", None) or default

    @staticmethod
    def _short_hash(obj, attr: str, default: str = "") -> str:
        """First 16 characters of obj.<attr>, or default if missing."""
        value = getattr(obj, attr, None) if obj is not None else None
        return str(value)[:16] if value else default

    def _approved_logs(self):
        """Approval logs with action='approved', ordered by level."""
        from app.models import ApprovalLog

        try:
            return list(
                self.requisition.approval_logs.filter_by(action="approved")
                .order_by(ApprovalLog.approval_level)
                .all()
            )
        except Exception:  # pragma: no cover - defensive
            current_app.logger.exception(
                "Could not fetch approved logs for requisition %s",
                getattr(self.requisition, "id", None),
            )
            return []

    def _get_security_info(self) -> dict[str, str]:
        """Document code, authorised amount and verification URL for the requisition."""
        from flask import url_for
        from app.models import ApprovalLog

        req = self.requisition

        document_code = ""
        try:
            from app.routes.requisition import _document_code

            approval_logs = (
                ApprovalLog.query.filter_by(requisition_id=req.id)
                .order_by(ApprovalLog.approval_level)
                .all()
            )
            document_code = _document_code(req, approval_logs)
        except Exception:  # pragma: no cover - defensive
            current_app.logger.exception(
                "Could not compute document code for requisition %s",
                getattr(req, "id", None),
            )

        try:
            verify_url = url_for(
                "requisition.verify", id=req.id, code=document_code, _external=True
            )
        except Exception:  # pragma: no cover - fallback URL
            verify_url = f"/verify/{req.id}"

        try:
            amount_value = float(getattr(req, "total_debit", 0) or 0)
            currency = getattr(req, "currency", "") or ""
            amount = f"{currency} {amount_value:.2f}".strip()
        except (TypeError, ValueError):
            amount = ""

        return {
            "document_code": document_code,
            "authorised_amount": amount,
            "verify_url": verify_url,
        }
