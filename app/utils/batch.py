"""Stanbic Business Online batch files: CSV upload and branded Excel copy."""
import csv
import io
import os
from decimal import Decimal, ROUND_HALF_UP

from flask import current_app

COMPANY_NAME = 'First Mutual Life Assurance Company'
BRAND_RED = 'D21F3C'
BRAND_BLACK = '1A1A1A'

# Column order required for the batch: Account Name, Account Number, Code, Amount, Reference, Currency Code, Branch Name
COLUMNS = ['Account Name', 'Account Number', 'Code', 'Amount', 'Reference', 'Currency Code', 'Branch Name']
AMOUNT_INDEX = 3


def money(value):
    """Round to two decimal places (half up), e.g. 1234.5 -> Decimal('1234.50')."""
    return Decimal(str(value or 0)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def payment_code(requisition):
    """'Code' column: the account code of the first line item on the requisition."""
    first = requisition.items.first()
    return first.account_code if first else ''


def batch_values(requisition):
    return [
        requisition.in_favor_of or '',
        requisition.account_number or '',
        payment_code(requisition),
        money(requisition.total_debit),
        requisition.requisition_number,
        requisition.currency or '',
        requisition.branch or '',
    ]


def to_batch_csv(requisitions):
    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow(COLUMNS)
    for r in requisitions:
        values = batch_values(r)
        values[AMOUNT_INDEX] = f"{values[AMOUNT_INDEX]:.2f}"
        writer.writerow(values)
    buf = io.BytesIO(text.getvalue().encode('utf-8'))
    buf.seek(0)
    return buf


def _logo_path():
    return os.path.join(current_app.root_path, 'static', 'images', 'logo.png')


def to_batch_excel(requisitions, batch_number, currency, stamp):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Batch'
    ws.sheet_view.showGridLines = False
    last_col = len(COLUMNS)

    # Logo sits in A1:B4
    try:
        with open(_logo_path(), 'rb') as fh:
            img = XLImage(io.BytesIO(fh.read()))
        ratio = img.width / img.height if img.height else 1
        img.height = 70
        img.width = int(70 * ratio)
        ws.add_image(img, 'A1')
    except Exception as exc:
        current_app.logger.warning('Batch Excel logo not added: %s', exc)

    for row, height in ((1, 28), (2, 20), (3, 16), (4, 16)):
        ws.row_dimensions[row].height = height

    banner = [
        (1, COMPANY_NAME, Font(bold=True, size=16, color=BRAND_RED)),
        (2, f'Payment batch {batch_number}', Font(bold=True, size=12)),
        (3, f'Currency: {currency}   |   Payments: {len(requisitions)}   |   Generated: {stamp}',
         Font(size=9, color='666666')),
    ]
    for row, text, font in banner:
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=last_col)
        cell = ws.cell(row=row, column=3, value=text)
        cell.font = font
        cell.alignment = Alignment(horizontal='left', vertical='center')

    red_rule = Side(style='medium', color=BRAND_RED)
    for col in range(1, last_col + 1):
        ws.cell(row=4, column=col).border = Border(bottom=red_rule)

    widths = [28, 18, 14, 16, 18, 14, 24]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    side = Side(style='thin', color='D0D0D0')
    border = Border(left=side, right=side, top=side, bottom=side)
    header_row = 6
    ws.row_dimensions[header_row].height = 26
    for col, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=BRAND_BLACK, end_color=BRAND_BLACK, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    zebra = PatternFill(start_color='F8F8F8', end_color='F8F8F8', fill_type='solid')
    first_data = header_row + 1
    row_num = first_data
    for idx, r in enumerate(requisitions, start=1):
        values = batch_values(r)
        for col, value in enumerate(values, start=1):
            if col - 1 == AMOUNT_INDEX:
                value = float(value)  # already rounded to 2 dp
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = border
            if idx % 2 == 0:
                cell.fill = zebra
            if col - 1 == AMOUNT_INDEX:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
        row_num += 1

    last_data = row_num - 1
    total_fill = PatternFill(start_color='FCE8EA', end_color='FCE8EA', fill_type='solid')
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.border = border
    ws.cell(row=row_num, column=1, value='TOTAL')
    amount_letter = get_column_letter(AMOUNT_INDEX + 1)
    total = ws.cell(row=row_num, column=AMOUNT_INDEX + 1,
                    value=f'=SUM({amount_letter}{first_data}:{amount_letter}{last_data})')
    total.number_format = '#,##0.00'
    total.alignment = Alignment(horizontal='right')

    ws.freeze_panes = f'A{first_data}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    from openpyxl.worksheet.properties import PageSetupProperties
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{header_row}:{header_row}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

