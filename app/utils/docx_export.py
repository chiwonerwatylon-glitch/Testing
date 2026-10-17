import io
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from flask import current_app

from app.utils.reports import AMOUNT_INDEX, BRAND_BLACK, BRAND_RED, COMPANY_NAME, build_rows, totals_by_currency

GREY = RGBColor(0x66, 0x66, 0x66)


def _logo():
    return os.path.join(current_app.root_path, 'static', 'images', 'logo.png')


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def _format_cell(cell, size, bold=False, color=None):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.bold = bold
            if color is not None:
                run.font.color.rgb = color


def to_docx(requisitions, meta):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.5)

    try:
        doc.add_picture(_logo(), width=Cm(3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as exc:
        current_app.logger.warning('Word report logo not added: %s', exc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(COMPANY_NAME.upper())
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor.from_string(BRAND_RED)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run('Requisition Report')
    sub_run.bold = True
    sub_run.font.size = Pt(12)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run(f"Generated: {meta['generated']}   |   Filters: {meta['filters']}")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = GREY

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_run = summary.add_run(f"Requisitions: {len(requisitions)}   |   Totals: {totals_by_currency(requisitions)}")
    s_run.bold = True
    s_run.font.size = Pt(10)

    headers, rows = build_rows(requisitions)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        _format_cell(cell, 8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(cell, BRAND_BLACK)

    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    header_props.append(repeat)

    if rows:
        for row in rows:
            cells = table.add_row().cells
            for idx, value in enumerate(row):
                cells[idx].text = str(value)
                _format_cell(cells[idx], 8)
    else:
        cells = table.add_row().cells
        cells[0].text = 'No requisitions match these filters.'
        _format_cell(cells[0], 8)

    total_cells = table.add_row().cells
    total_cells[0].text = 'TOTAL'
    total_cells[AMOUNT_INDEX].text = totals_by_currency(requisitions)
    for cell in total_cells:
        _shade(cell, 'FCE8EA')
        _format_cell(cell, 8, bold=True)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f_run = footer.add_run(f"Confidential. Generated {meta['generated']}.")
    f_run.font.size = Pt(8)
    f_run.font.color.rgb = GREY

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
